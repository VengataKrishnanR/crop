"""Review downloads — the request and response of every UAT test, per test or per run.

Reviewers see results on screen; the evidence (exact request sent, exact
response received, and the second environment's response and differences when
a comparison ran) is downloaded rather than shown.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from tools.core.db import get_session
from tools.core.models import OnboardingCase, UatComparison, UatResult, UatRun, UatTestCase
from tools.workflow import State

from .auth import CurrentUser, current_user, ensure_country

router = APIRouter()

_BUSINESS_STATES = {State.UAT_SIGNOFF, State.RTP, State.COMPLETED}


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name)[:80] or "file"


def _json_bytes(obj: Any) -> bytes:
    if isinstance(obj, (bytes, bytearray)):
        return bytes(obj)
    if isinstance(obj, str):
        try:
            obj = json.loads(obj)
        except json.JSONDecodeError:
            return obj.encode("utf-8")
    return json.dumps(obj, indent=2, ensure_ascii=False, default=str).encode("utf-8")


def _authorised_run(session: Session, run_id: str, user: CurrentUser) -> tuple[UatRun, OnboardingCase]:
    run = session.get(UatRun, run_id)
    if run is None:
        raise HTTPException(404, "No such run.")
    case = session.get(OnboardingCase, run.case_id)
    ensure_country(user, case.ctry_cd)
    if user.is_business and State(case.state) not in _BUSINESS_STATES:
        raise HTTPException(404, "No such run.")
    return run, case


def _attachment(content: bytes, filename: str, media: str) -> Response:
    return Response(content, media_type=media,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/uat/results/{result_id}/download", tags=["uat"])
def download_result(result_id: str, part: str = "response", user: CurrentUser = Depends(current_user),
                    session: Session = Depends(get_session)) -> Response:
    """``part`` = ``request`` | ``response`` | ``compare`` (second environment + differences)."""
    result = session.get(UatResult, result_id)
    if result is None:
        raise HTTPException(404, "No such result.")
    _authorised_run(session, result.run_id, user)
    name = _safe(result.test_id)
    if part == "request":
        return _attachment(_json_bytes(result.payload), f"{name}_request.json", "application/json")
    if part == "response":
        body = result.response_body if result.response_body is not None else {"error": result.error}
        return _attachment(_json_bytes(body), f"{name}_response.json", "application/json")
    if part == "compare":
        cmp_ = session.execute(select(UatComparison).where(UatComparison.result_id == result.id)
                               ).scalar_one_or_none()
        if cmp_ is None:
            raise HTTPException(404, "No comparison was run for this test.")
        doc = {"endpoint": cmp_.endpoint_name, "status": cmp_.response_status, "identical": cmp_.identical,
               "summary": cmp_.summary, "differences": cmp_.differences, "error": cmp_.error,
               "response": cmp_.response_body}
        return _attachment(_json_bytes(doc), f"{name}_compare_{_safe(cmp_.endpoint_name)}.json",
                           "application/json")
    raise HTTPException(400, "part must be request, response or compare.")


def _summary_xlsx(case: OnboardingCase, run: UatRun, rows: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append([f"UAT run — {case.code} ({case.ctry_cd}) — cycle {run.cycle} — {run.executor}"])
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="D40511")
    ws.append([f"Passed {run.passed} · Failed {run.failed} · Errors {run.errors} · "
               f"Started {run.started_at:%Y-%m-%d %H:%M} by {run.started_by}"])
    ws.append([])
    headers = ["Test", "What is tested", "Result", "Reason", "HTTP status", "Latency (ms)",
               "Compared with", "Identical", "Differences", "Request file", "Response file"]
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2D2D2D")
    fills = {"PASS": "D4EDDA", "FAIL": "F8D7DA", "ERROR": "FDE8D8", "INCONCLUSIVE": "FFF3CD"}
    for r in rows:
        ws.append([r["test_id"], r["summary"], r["verdict"], r["reason"], r["status"], r["latency"],
                   r["compare_with"], r["identical"], r["differences"], r["request_file"], r["response_file"]])
        ws.cell(row=ws.max_row, column=3).fill = PatternFill("solid", fgColor=fills.get(r["verdict"], "FFFFFF"))
    for col, width in zip("ABCDEFGHIJK", (16, 50, 12, 60, 11, 12, 18, 10, 12, 30, 30)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/uat/runs/{run_id}/download", tags=["uat"])
def download_run(run_id: str, user: CurrentUser = Depends(current_user),
                 session: Session = Depends(get_session)) -> Response:
    """A zip: ``summary.xlsx``, ``requests/``, ``responses/`` and (if compared) ``compare/``."""
    run, case = _authorised_run(session, run_id, user)
    results = list(session.execute(select(UatResult).where(UatResult.run_id == run.id)
                                   .order_by(UatResult.test_id)).scalars())
    cases = {tc.id: tc for tc in session.execute(
        select(UatTestCase).where(UatTestCase.case_id == run.case_id)).scalars()}
    comparisons = {c.result_id: c for c in session.execute(
        select(UatComparison).where(UatComparison.result_id.in_([r.id for r in results]))).scalars()}

    buf = io.BytesIO()
    rows = []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for r in results:
            name = _safe(r.test_id)
            req_file, resp_file = f"requests/{name}.json", f"responses/{name}.json"
            zf.writestr(req_file, _json_bytes(r.payload))
            zf.writestr(resp_file, _json_bytes(r.response_body if r.response_body is not None
                                               else {"error": r.error}))
            cmp_ = comparisons.get(r.id)
            if cmp_ is not None:
                zf.writestr(f"compare/{name}_{_safe(cmp_.endpoint_name)}_response.json",
                            _json_bytes(cmp_.response_body if cmp_.response_body is not None
                                        else {"error": cmp_.error}))
                zf.writestr(f"compare/{name}_differences.json",
                            _json_bytes({"identical": cmp_.identical, "summary": cmp_.summary,
                                         "differences": cmp_.differences}))
            tc = cases.get(r.test_case_id)
            rows.append({
                "test_id": r.test_id, "summary": (tc.business_summary or tc.scenario) if tc else "",
                "verdict": r.verdict,
                "reason": r.error or "; ".join((r.check or {}).get("reasons") or []),
                "status": r.response_status, "latency": r.latency_ms,
                "compare_with": cmp_.endpoint_name if cmp_ else "",
                "identical": ("yes" if cmp_.identical else "no") if cmp_ else "",
                "differences": (cmp_.summary or {}).get("totalDifferences") if cmp_ else "",
                "request_file": req_file, "response_file": resp_file,
            })
        zf.writestr("summary.xlsx", _summary_xlsx(case, run, rows))
    filename = f"{_safe(case.code)}_UAT_cycle{run.cycle}.zip"
    return _attachment(buf.getvalue(), filename, "application/zip")
