"""Rule configuration: the onboarding's M2C draft edited row by row, its bindings,
the condition builder, changes vs the workbook, package coverage and the
generated M2C workbook — plus who submits the requirement and the timeline."""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from conftest import FULL_M2C
from openpyxl import load_workbook

from test_v3_collab import BUSINESS, CONSULTANT, PM, TEAM, _to_rule_config, loaded  # noqa: F401

pytestmark = pytest.mark.workflow


def _ws(client: TestClient, case_id: str, headers=TEAM) -> dict:
    resp = client.get(f"/api/onboardings/{case_id}/rule-config", headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


# --------------------------------------------------------------------------- #
# who submits
# --------------------------------------------------------------------------- #

def test_only_the_business_submits_its_requirement(client: TestClient) -> None:
    case_id = client.post("/api/onboardings", json={"ctry_cd": "DE", "title": "t"}, headers=BUSINESS).json()["id"]
    client.post(f"/api/onboardings/{case_id}/intake/form", headers=BUSINESS,
                json={"m2c": FULL_M2C, "description": "DE needs POA when DDP", "requirement_groups": ["POA"]})
    consultant = client.get(f"/api/onboardings/{case_id}", headers=CONSULTANT).json()
    assert not any(a["name"] == "submit_requirement" and a["mine"] for a in consultant["actions"])
    assert client.post(f"/api/onboardings/{case_id}/transition", json={"name": "submit_requirement"},
                       headers=CONSULTANT).status_code in (403, 409)
    business = client.get(f"/api/onboardings/{case_id}", headers=BUSINESS).json()
    assert any(a["name"] == "submit_requirement" and a["mine"] for a in business["actions"])


def test_the_pm_owns_the_timeline(client: TestClient) -> None:
    case_id = client.post("/api/onboardings", json={"ctry_cd": "DE", "title": "t"}, headers=BUSINESS).json()["id"]
    client.post(f"/api/onboardings/{case_id}/intake/form", headers=BUSINESS,
                json={"m2c": FULL_M2C, "description": "DE needs POA when DDP", "requirement_groups": ["POA"]})
    for name, who in (("submit_requirement", BUSINESS), ("approve_requirement", CONSULTANT)):
        assert client.post(f"/api/onboardings/{case_id}/transition", json={"name": name}, headers=who).status_code == 200
    # Waiting for the PM, not the consultant.
    view = client.get(f"/api/onboardings/{case_id}", headers=CONSULTANT).json()
    assert view["waiting"]["roles"] == ["PM"] and not any(a["mine"] for a in view["actions"]
                                                          if a["name"] != "withdraw")
    answers = {"answers": {"requested_rtp": "2027-03-31", "earliest_start": "2026-11-02"}}
    assert client.post(f"/api/onboardings/{case_id}/timeline/generate", headers=CONSULTANT,
                       json=answers).status_code == 403
    assert client.post(f"/api/onboardings/{case_id}/timeline/generate", headers=PM, json=answers).status_code == 200
    pm_view = client.get(f"/api/onboardings/{case_id}", headers=PM).json()
    assert "approve_timeline" in {a["name"] for a in pm_view["actions"] if a["mine"]}
    assert client.post(f"/api/onboardings/{case_id}/transition", json={"name": "approve_timeline"},
                       headers=CONSULTANT).status_code in (403, 409)
    assert client.post(f"/api/onboardings/{case_id}/transition", json={"name": "approve_timeline"},
                       headers=PM).status_code == 200
    assert client.get(f"/api/onboardings/{case_id}", headers=PM).json()["state"] == "RULE_CONFIG"
    assert client.get(f"/api/onboardings/{case_id}/timeline", headers=BUSINESS).json()["status"] == "APPROVED"


# --------------------------------------------------------------------------- #
# rows and bindings
# --------------------------------------------------------------------------- #

def test_workspace_shows_rows_with_their_checks(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "AE")
    ws = _ws(client, case_id)
    assert ws["editable"] is True and ws["summary"]["rows"] == 3
    row52 = next(r for r in ws["rows"] if r["row_no"] == 52)
    assert row52["status"] == "error"
    assert any(i["field"] == "expiry_date" for i in row52["issues"])
    assert row52["document_reference"] == "AE|Import||"
    # Business users never see the M2C workspace.
    assert client.get(f"/api/onboardings/{case_id}/rule-config", headers=BUSINESS).status_code == 403


def test_errors_block_submitting_rule_config_until_fixed(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "AE")
    client.post(f"/api/onboardings/{case_id}/dmg/not-required", json={"comment": "No DMG change"}, headers=CONSULTANT)
    overview = client.get(f"/api/onboardings/{case_id}", headers=TEAM).json()
    submit = next(a for a in overview["actions"] if a["name"] == "request_sanity_check")
    assert not submit["allowed"] and any("M2C row" in b for b in submit["blocked_by"])

    row52 = next(r for r in _ws(client, case_id)["rows"] if r["row_no"] == 52)
    fixed = client.put(f"/api/onboardings/{case_id}/rule-config/rows/{row52['id']}", headers=TEAM,
                       json={"expiry_date": "2099-12-31"})
    assert fixed.status_code == 200, fixed.text
    assert not any(i["level"] == "error" for i in fixed.json()["issues"])
    overview = client.get(f"/api/onboardings/{case_id}", headers=TEAM).json()
    submit = next(a for a in overview["actions"] if a["name"] == "request_sanity_check")
    assert not any("M2C row" in b for b in submit["blocked_by"])


def test_add_edit_duplicate_delete_row_and_bindings(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "DE")
    ws = _ws(client, case_id)
    group = ws["options"]["groups_mapped_for_country"][0]
    bad = client.post(f"/api/onboardings/{case_id}/rule-config/rows", headers=TEAM,
                      json={"shipment_type": "B2C", "min_value": 500, "max_value": 100, "sirg_cd": group})
    assert bad.status_code == 200
    assert any(i["field"] == "max_value" and i["level"] == "error" for i in bad.json()["issues"])
    row = client.put(f"/api/onboardings/{case_id}/rule-config/rows/{bad.json()['id']}", headers=TEAM,
                     json={"min_value": 0, "max_value": 150, "currency": "eur",
                           "other_logic": '!serviceCodeList.contains("DD")'}).json()
    assert row["currency"] == "EUR" and row["status"] != "error"

    defs = client.get(f"/api/onboardings/{case_id}/rule-config/definitions", params={"group": group},
                      headers=TEAM).json()
    sir = next(d for d in defs if d["mapped_in_group"])
    attrs = client.get(f"/api/onboardings/{case_id}/rule-config/definitions/{sir['sir_id']}/attributes",
                       headers=TEAM).json()
    assert attrs
    bound = client.post(f"/api/onboardings/{case_id}/rule-config/rows/{row['id']}/bindings", headers=TEAM,
                        json={"sir_id": sir["sir_id"], "waived_to_validate": True,
                              "attribute_ids": [attrs[0]["att_id"]]})
    assert bound.status_code == 200, bound.text
    b = bound.json()["bindings"][0]
    assert b["cmpt_nm"] == sir["cmpt_nm"] and b["waived_to_validate"] is True
    assert not any(i["field"] == "bindings" for i in bound.json()["issues"])
    dup = client.post(f"/api/onboardings/{case_id}/rule-config/rows/{row['id']}/bindings", headers=TEAM,
                      json={"sir_id": sir["sir_id"]})
    assert dup.status_code == 409
    unknown_attr = client.post(f"/api/onboardings/{case_id}/rule-config/rows/{row['id']}/bindings", headers=TEAM,
                               json={"sir_id": defs[-1]["sir_id"], "attribute_ids": ["nope"]})
    assert unknown_attr.status_code in (400, 409)

    copy = client.post(f"/api/onboardings/{case_id}/rule-config/rows/{row['id']}/duplicate", headers=TEAM).json()
    assert copy["id"] != row["id"] and len(copy["bindings"]) == 1
    assert client.delete(f"/api/onboardings/{case_id}/rule-config/rows/{copy['id']}/bindings/{copy['bindings'][0]['id']}",
                         headers=TEAM).json()["bindings"] == []
    assert client.delete(f"/api/onboardings/{case_id}/rule-config/rows/{copy['id']}", headers=TEAM).status_code == 200
    assert _ws(client, case_id)["summary"]["rows"] == 2

    # Consultants may edit too; PMs and the business may not.
    assert client.put(f"/api/onboardings/{case_id}/rule-config/rows/{row['id']}", headers=PM,
                      json={"currency": "USD"}).status_code == 403
    audit = client.get(f"/api/onboardings/{case_id}/audit", headers=TEAM)
    if audit.status_code == 200:
        assert any(e["action"] == "m2c:row_edited" for e in audit.json())


def test_rows_are_read_only_outside_rule_config(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "DE")
    client.post(f"/api/onboardings/{case_id}/dmg/not-required", json={"comment": "No DMG change"}, headers=CONSULTANT)
    assert client.post(f"/api/onboardings/{case_id}/transition", json={"name": "request_sanity_check"},
                       headers=TEAM).status_code == 200
    ws = _ws(client, case_id)
    assert ws["editable"] is False
    resp = client.put(f"/api/onboardings/{case_id}/rule-config/rows/{ws['rows'][0]['id']}", headers=TEAM,
                      json={"currency": "USD"})
    assert resp.status_code == 409


# --------------------------------------------------------------------------- #
# condition builder, changes, coverage, workbook
# --------------------------------------------------------------------------- #

def test_condition_builder_writes_valid_other_logic(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "BG")
    built = client.post(f"/api/onboardings/{case_id}/rule-config/condition", headers=TEAM, json={
        "clauses": [{"type": "short_code", "op": "is present (has a value)", "values": "SDT"},
                    {"type": "service_code", "op": "has none of", "values": "DD"}]}).json()
    # Exactly the business-validated BG IOSS logic.
    assert built["expression"] == 'matchCiMask("shortCode|==|SDT","value|!=|<blank>") && !serviceCodeList.contains("DD")'
    assert built["parses"] and not built["errors"]
    multi = client.post(f"/api/onboardings/{case_id}/rule-config/condition", headers=TEAM, json={
        "joiner": "and", "clauses": [{"type": "origin_country", "op": "is one of", "values": "CN, US"},
                                     {"type": "incoterm", "op": "is one of", "values": ["DDP"]}]}).json()
    assert multi["expression"].startswith('(OriginCountry=="CN" || OriginCountry=="US") && ')
    bad = client.post(f"/api/onboardings/{case_id}/rule-config/condition", headers=TEAM,
                      json={"clauses": [{"type": "short_code", "op": "is blank", "values": ""}]})
    assert bad.status_code == 400
    palette = client.get(f"/api/onboardings/{case_id}/rule-config/palette", headers=TEAM).json()
    assert palette and all("ok" in p for p in palette)


def test_changes_coverage_and_m2c_workbook(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "BG")
    ws = _ws(client, case_id)
    rows = {r["row_no"]: r for r in ws["rows"]}
    assert rows[56]["status"] == "package"          # BG is not in the config package yet
    client.put(f"/api/onboardings/{case_id}/rule-config/rows/{rows[57]['id']}", headers=TEAM,
               json={"other_logic": 'matchCiMask("shortCode|==|SDT","value|!=|<blank>") && !serviceCodeList.contains("DD")'})
    client.post(f"/api/onboardings/{case_id}/rule-config/rows", headers=TEAM,
                json={"shipment_type": "B2C", "currency": "EUR", "max_value": 150, "sirg_cd": "AUTH"})

    ch = client.get(f"/api/onboardings/{case_id}/rule-config/changes", headers=TEAM).json()
    assert ch["counts"] == {"NEW": 1, "CHANGED": 1, "REMOVED": 0, "UNCHANGED": 1}
    changed = next(i for i in ch["items"] if i["status"] == "CHANGED")
    assert changed["row_no"] == 57 and set(changed["fields"]) == {"other_logic"}

    cov = client.get(f"/api/onboardings/{case_id}/rule-config/coverage", headers=TEAM).json()
    assert cov["complete"] is False and cov["support_request"]
    assert all(r["CtryProc.DhlCtryCd"] == "BG" for r in cov["support_request"])
    csv_text = client.get(f"/api/onboardings/{case_id}/rule-config/coverage.csv", headers=TEAM).text
    assert csv_text.startswith("CtryProc.DhlCtryCd,CtryProc.ProcLvlCd")

    resp = client.get(f"/api/onboardings/{case_id}/rule-config/m2c.xlsx", headers=TEAM)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert {"Import", "Export", "SIR Reference", "Technical SIR", "Technical SIRG", "Other Logic",
            "Validation", "Changes (platform)"} <= set(wb.sheetnames)
    imp = wb["Import"]
    assert imp.max_row == 4                          # header + 3 BG rows
    assert imp["P2"].value == '=CONCATENATE(B2,"|Import|",N2,"|",O2)'
    ref = wb["SIR Reference"]
    assert ref.max_row >= 3 and ref["A2"].value.startswith("=IF(Import!P")
    log = [row for row in wb["Changes (platform)"].iter_rows(min_row=2, values_only=True)]
    assert ("CHANGED", 57, "Other Logic") in [r[:3] for r in log]
