"""The guided sanity check: six checks, saved runs with a fingerprint, waivers and the report."""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from test_e2e_positive import align_rows
from test_v3_collab import BUSINESS, CONSULTANT, TEAM, _to_rule_config, loaded  # noqa: F401

pytestmark = pytest.mark.workflow


def _sanity(client: TestClient, case_id: str) -> dict:
    resp = client.get(f"/api/onboardings/{case_id}/sanity", headers=TEAM)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _to_sanity(client: TestClient) -> str:
    case_id = _to_rule_config(client, "DE")
    align_rows(client, case_id)
    client.post(f"/api/onboardings/{case_id}/dmg/not-required", headers=CONSULTANT, json={"comment": "none"})
    assert client.post(f"/api/onboardings/{case_id}/transition", json={"name": "request_sanity_check"},
                       headers=TEAM).status_code == 200
    return case_id


def test_entering_the_stage_runs_it_and_a_clean_run_can_be_approved(client: TestClient, loaded: None) -> None:
    case_id = _to_sanity(client)
    s = _sanity(client, case_id)
    assert s["ready"] and s["latest_run"]["status"] == "PASS" and s["latest_run"]["current"]
    assert s["latest_run"]["run_by"] == "system:on-entry"
    assert client.post(f"/api/onboardings/{case_id}/transition", json={"name": "sanity_passed"},
                       headers=TEAM).status_code == 200


def test_a_change_after_the_run_needs_a_new_run(client: TestClient, loaded: None) -> None:
    case_id = _to_sanity(client)
    # Back to rule config, change a row, return: the earlier passing run no longer applies.
    assert client.post(f"/api/onboardings/{case_id}/transition", headers=TEAM, json={
        "name": "sanity_failed", "comment": "Tweak the currency"}).status_code == 200
    row = client.get(f"/api/onboardings/{case_id}/rule-config", headers=TEAM).json()["rows"][0]
    client.put(f"/api/onboardings/{case_id}/rule-config/rows/{row['id']}", headers=TEAM, json={"currency": "USD"})
    s = _sanity(client, case_id)
    assert s["latest_run"]["current"] is False
    # The row still matches nothing else, but the requirement does not mention currency -> differs now.
    rm = next(i for i in s["items"] if i["key"] == "requirement_match")
    assert rm["status"] == "fail" and any(d["field"] == "Currency" for d in rm["findings"][0]["differences"])


def test_send_back_reason_is_written_from_the_failing_checks(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "DE")
    client.post(f"/api/onboardings/{case_id}/dmg/not-required", headers=CONSULTANT, json={"comment": "none"})
    client.post(f"/api/onboardings/{case_id}/transition", json={"name": "request_sanity_check"}, headers=TEAM)
    s = _sanity(client, case_id)
    assert not s["ready"]
    assert "match the approved requirement" in s["send_back_reason"]
    assert "row 38 (POA) differs in Effective date" in s["send_back_reason"]
    item = next(i for i in s["items"] if i["key"] == "requirement_match")
    assert item["fix"]["action"] == "send_back" and item["waivable"]
    # Approving is refused with the reason; sending back with the prefilled text works.
    resp = client.post(f"/api/onboardings/{case_id}/transition", json={"name": "sanity_passed"}, headers=TEAM)
    assert resp.status_code == 409 and "failed" in " ".join(resp.json()["detail"]["reasons"])
    assert client.post(f"/api/onboardings/{case_id}/transition", headers=TEAM, json={
        "name": "sanity_failed", "comment": s["send_back_reason"]}).status_code == 200


def test_a_deliberate_difference_can_be_accepted_and_revoked(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "DE")
    client.post(f"/api/onboardings/{case_id}/dmg/not-required", headers=CONSULTANT, json={"comment": "none"})
    client.post(f"/api/onboardings/{case_id}/transition", json={"name": "request_sanity_check"}, headers=TEAM)
    w = client.post(f"/api/onboardings/{case_id}/sanity/waivers", headers=CONSULTANT, json={
        "item_key": "requirement_match", "reason": "Row 38 is the live DE row; requirement only adds POA scope"})
    assert w.status_code == 200, w.text
    s = _sanity(client, case_id)
    assert s["ready"] and s["latest_run"]["status"] == "PASS" and s["latest_run"]["current"]
    assert client.post(f"/api/onboardings/{case_id}/sanity/waivers", headers=CONSULTANT, json={
        "item_key": "requirement_match", "reason": "again please"}).status_code == 409   # nothing failing now
    assert client.delete(f"/api/onboardings/{case_id}/sanity/waivers/{w.json()['id']}",
                         headers=CONSULTANT).status_code == 200
    s = _sanity(client, case_id)
    assert not s["ready"]
    assert any(x["revoked_by"] for x in s["waivers"])                  # kept, marked revoked


def test_run_rules_and_report(client: TestClient, loaded: None) -> None:
    case_id = _to_rule_config(client, "DE")
    # Before the stage: a preview only.
    assert client.post(f"/api/onboardings/{case_id}/sanity/run", headers=TEAM).status_code == 409
    assert _sanity(client, case_id)["can_run"] is False
    case_id = _to_sanity(client)
    assert client.post(f"/api/onboardings/{case_id}/sanity/run", headers=BUSINESS).status_code == 403
    r = client.post(f"/api/onboardings/{case_id}/sanity/run", headers=TEAM).json()
    assert r["status"] == "PASS" and r["current"]
    rep = client.get(f"/api/onboardings/{case_id}/sanity/report.xlsx", headers=TEAM)
    wb = load_workbook(io.BytesIO(rep.content))
    assert wb.sheetnames == ["Summary", "Findings"]
    titles = [row[0] for row in wb["Summary"].iter_rows(values_only=True)]
    assert "The config package can raise every requirement" in titles
