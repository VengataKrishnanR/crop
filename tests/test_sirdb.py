"""SIR/SIRG reference database — schema, loader, resolution and the final Import output.

Runs against a real PostgreSQL (SIRDB_DATABASE_URL or deploy/sirdb.env). Each test session
creates its own throw-away database and drops it afterwards; the loaded data is never touched.
Skipped when no PostgreSQL or no M2C workbook is available.
"""
from __future__ import annotations

import shutil
import uuid
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

import openpyxl
import pytest

psycopg = pytest.importorskip("psycopg")

from conftest import REFERENCE  # noqa: E402
from tools.sirdb import loader, migrate, resolve  # noqa: E402
from tools.sirdb.db import connect, database_url  # noqa: E402
from tools.sirdb.profile import profile  # noqa: E402
from tools.sirdb.workbook import read_workbook  # noqa: E402


def _admin_url() -> str:
    try:
        return database_url()
    except RuntimeError:
        pytest.skip("SIRDB_DATABASE_URL not configured")


@pytest.fixture(scope="module")
def db_url():
    admin = _admin_url()
    name = f"sirdb_test_{uuid.uuid4().hex[:8]}"
    try:
        with psycopg.connect(admin, autocommit=True) as c:
            c.execute(f'CREATE DATABASE "{name}"')
    except psycopg.OperationalError:
        pytest.skip("PostgreSQL not reachable")
    p = urlsplit(admin)
    url = urlunsplit((p.scheme, p.netloc, f"/{name}", p.query, p.fragment))
    try:
        yield url
    finally:
        with psycopg.connect(admin, autocommit=True) as c:
            c.execute(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)')


@pytest.fixture(scope="module")
def workbook(tmp_path_factory) -> Path:
    src = sorted(REFERENCE.glob("M2c*.xlsx"))
    if not src:
        pytest.skip("no M2C workbook in Reference/")
    dst = tmp_path_factory.mktemp("wb") / "m2c.xlsx"
    shutil.copyfile(src[0], dst)          # never open the original for writing
    return dst


@pytest.fixture(scope="module")
def loaded(db_url, workbook):
    assert migrate.migrate(db_url)                      # all five applied
    res = loader.load(workbook, db_url, actor="test")
    assert res["status"] == "LOADED"
    return res


def q(url, sql, *args):
    with connect(url, actor="test") as c, c.cursor() as cur:
        cur.execute(sql, args)
        return cur.fetchall()


# --------------------------------------------------------------------------- #
# schema
# --------------------------------------------------------------------------- #

def test_migrations_are_versioned_and_rerunnable(db_url, loaded) -> None:
    assert migrate.migrate(db_url) == []                # nothing pending
    versions = [r[0] for r in q(db_url, "SELECT version FROM sirdb.schema_migration ORDER BY version")]
    assert versions == ["0001", "0002", "0003", "0004", "0005"]
    views = {r[0] for r in q(db_url, "SELECT table_name FROM information_schema.views WHERE table_schema='sirdb'")}
    assert {"v_import_final", "v_import_rule_sirg", "v_import_rule_sir_reference"} <= views


def test_changed_migration_is_refused(db_url, loaded, monkeypatch, tmp_path) -> None:
    for f in migrate.migration_files():
        (tmp_path / f.name).write_text(f.read_text(encoding="utf-8") + "\n-- edited", encoding="utf-8")
    monkeypatch.setattr(migrate, "MIGRATIONS", tmp_path)
    with pytest.raises(RuntimeError, match="changed after it was applied"):
        migrate.migrate(db_url)


# --------------------------------------------------------------------------- #
# loading: lossless, idempotent, traceable
# --------------------------------------------------------------------------- #

def test_every_cell_and_row_is_preserved(db_url, loaded, workbook) -> None:
    wb = read_workbook(workbook)
    for s in wb.sheets.values():
        cells = sum(len(v) for v in s.rows.values())
        assert q(db_url, "SELECT count(*) FROM sirdb.source_cell WHERE batch_id=%s AND sheet_name=%s",
                 loaded["batch_id"], s.name)[0][0] == cells, s.name
        formulas = sum(1 for v in s.rows.values() for c in v if c.formula)
        assert q(db_url, "SELECT count(*) FROM sirdb.source_cell WHERE batch_id=%s AND sheet_name=%s "
                         "AND formula IS NOT NULL", loaded["batch_id"], s.name)[0][0] == formulas
    unmapped = q(db_url, "SELECT sheet_name, row_no FROM sirdb.source_row WHERE batch_id=%s AND NOT is_header "
                         "AND map_status = 'UNMAPPED'", loaded["batch_id"])
    assert unmapped == []
    # Formula text and cached value side by side.
    f, v = q(db_url, "SELECT formula, value_text FROM sirdb.source_cell WHERE batch_id=%s AND sheet_name='Import' "
                     "AND row_no=2 AND col_index=16", loaded["batch_id"])[0]
    assert f == '=CONCATENATE(B2,"|Import|",N2,"|",O2)' and v == "ES|Import|DH|CIDM"


def test_blank_and_unheaded_columns_are_kept(db_url, loaded) -> None:
    hdr, name = q(db_url, "SELECT header_text, provisional_name FROM sirdb.source_column WHERE batch_id=%s "
                          "AND sheet_name='SIR Reference' AND col_index=9", loaded["batch_id"])[0]
    assert hdr is None and name == "col_09_unheaded"
    assert q(db_url, "SELECT count(*) FROM sirdb.sir_reference WHERE batch_id=%s AND col_09_unheaded IS NOT NULL",
             loaded["batch_id"])[0][0] == 49
    # SIRG Dsc is empty on every row and still a column, row count intact.
    assert q(db_url, "SELECT count(*), count(sirg_dsc) FROM sirdb.technical_sirg WHERE batch_id=%s",
             loaded["batch_id"])[0] == (1004, 0)
    # Headerless sheets keep their first row as data.
    assert q(db_url, "SELECT value_raw FROM sirdb.reference_list_item i JOIN sirdb.reference_list l USING (list_id) "
                     "WHERE l.batch_id=%s AND l.list_name='Sheet2' AND position=1", loaded["batch_id"])[0][0] == "ES"
    fields = q(db_url, "SELECT fields FROM sirdb.inferred_record WHERE batch_id=%s", loaded["batch_id"])[0][0]
    assert fields["cmpt_nm"] == "DEPOA" and fields["sirg_dsc"] is None


def test_source_values_are_never_repaired(db_url, loaded) -> None:
    raw, norm, why = q(db_url, "SELECT value_raw, value_normalized, normalization FROM sirdb.reference_list_item i "
                               "JOIN sirdb.reference_list l USING (list_id) WHERE l.batch_id=%s "
                               "AND value_raw LIKE 'SK%%'", loaded["batch_id"])[0]
    assert (raw, norm) == ("SK ", "SK") and "trimmed" in why
    raw, d, issue = q(db_url, "SELECT effective_date_raw, effective_date, effective_date_issue FROM sirdb.import_rule "
                              "WHERE batch_id=%s AND source_row_no=25", loaded["batch_id"])[0]
    assert raw == "28-02-2022" and str(d) == "2022-02-28" and issue is None
    raw, flag = q(db_url, "SELECT waived_to_validate_raw, waived_to_validate FROM sirdb.sir_reference "
                          "WHERE batch_id=%s AND source_row_no=82", loaded["batch_id"])[0]
    assert raw == "4" and flag is None                      # kept as written, not coerced
    assert q(db_url, "SELECT country_hold_code FROM sirdb.import_rule WHERE batch_id=%s AND source_row_no=5",
             loaded["batch_id"])[0][0] == "3"


def test_loading_the_same_file_again_changes_nothing(db_url, loaded, workbook) -> None:
    before = q(db_url, "SELECT (SELECT count(*) FROM sirdb.source_cell), (SELECT count(*) FROM sirdb.technical_sir), "
                       "(SELECT count(*) FROM sirdb.import_batch)")[0]
    again = loader.load(workbook, db_url, actor="test")
    assert again["status"] == "ALREADY_LOADED" and again["batch_id"] == loaded["batch_id"]
    assert q(db_url, "SELECT (SELECT count(*) FROM sirdb.source_cell), (SELECT count(*) FROM sirdb.technical_sir), "
                     "(SELECT count(*) FROM sirdb.import_batch)")[0] == before


def test_findings_are_stored_with_source_rows(db_url, loaded) -> None:
    rows = {r[0] for r in q(db_url, "SELECT row_no FROM sirdb.load_issue WHERE batch_id=%s "
                                    "AND code='IMP_EXPIRY_BEFORE_EFFECTIVE'", loaded["batch_id"])}
    assert rows == {25, 52, 54}
    placeholder_rows = {r[0] for r in q(db_url, "SELECT row_no FROM sirdb.load_issue WHERE batch_id=%s "
                                                "AND code='IMP_DOCREF_PLACEHOLDER'", loaded["batch_id"])}
    # 35-55: typed TEST|TEST|TEST; 32-33: the formula result carries the TBD hold code.
    assert placeholder_rows == set(range(35, 56)) | {32, 33}
    assert q(db_url, "SELECT row_no FROM sirdb.load_issue WHERE batch_id=%s AND code='TSIRG_EXACT_DUPLICATE'",
             loaded["batch_id"])[0][0] == 880


def test_source_rows_are_append_only_and_changes_are_logged(db_url, loaded) -> None:
    with connect(db_url, actor="test") as c, c.cursor() as cur:
        with pytest.raises(psycopg.errors.RaiseException, match="append-only"):
            cur.execute("UPDATE sirdb.source_row SET record = '[]'::jsonb WHERE batch_id=%s AND sheet_name='Import' "
                        "AND row_no=2", (loaded["batch_id"],))
    assert q(db_url, "SELECT count(*) FROM sirdb.change_log WHERE table_name='technical_sir' AND operation='INSERT'")[0][0] == 984


# --------------------------------------------------------------------------- #
# resolution
# --------------------------------------------------------------------------- #

def test_known_mappings_resolve(db_url, loaded) -> None:
    with connect(db_url) as c, c.cursor() as cur:
        de = resolve.resolve_sirg(cur, "DE", "POA")
        assert (de.status, de.cmpt_nm, de.sir_in_catalog) == ("RESOLVED", "DEPOA", True)
        assert de.sir_technical_id.startswith("b5a141ed")
        bg = resolve.resolve_sirg(cur, "BG", "POAB")
        assert (bg.status, bg.cmpt_nm) == ("RESOLVED", "POABBG")
        lines = resolve.sir_reference_lines(cur, "DE", "POA")
        assert [(l["source_row"], l["col_09_unheaded"]) for l in lines] == [(84, "Power of Attorney")]


def test_one_rule_can_have_many_sir_reference_lines(db_url, loaded) -> None:
    with connect(db_url) as c, c.cursor() as cur:
        lines = resolve.sir_reference_lines(cur, "MX", "IRID")
    assert [l["source_row"] for l in lines] == [67, 68, 69, 70]
    assert all(l["join"] == "SIRG_DOCUMENT_REFERENCE" for l in lines)


def test_missing_and_duplicate_mappings_are_reported_not_guessed(db_url, loaded) -> None:
    with connect(db_url) as c, c.cursor() as cur:
        missing = resolve.resolve_sirg(cur, "DE", "NOPE")
        assert missing.status == "NOT_FOUND" and "maintenance" in missing.message and missing.cmpt_nm is None
        dup = resolve.resolve_sirg(cur, "AU", "GSTE")
        assert dup.status == "DUPLICATE_RECORDS" and [m["source_row"] for m in dup.matches] == [877, 880]
    statuses = dict(q(db_url, "SELECT sirg_status, count(*) FROM sirdb.v_import_rule_sirg WHERE batch_id=%s "
                              "GROUP BY 1", loaded["batch_id"]))
    assert statuses == {"RESOLVED": 32, "NOT_APPLICABLE": 24}


def test_ambiguous_mapping_in_a_new_workbook_version(db_url, loaded, workbook, tmp_path) -> None:
    """A newer workbook with a conflicting SIRG row: new batch, becomes current, resolution says AMBIGUOUS."""
    wb = openpyxl.load_workbook(workbook)
    ws = wb["Technical SIRG"]
    de = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[4] == "DE" and r[1] == "POA")
    other = next(r for r in wb["Technical SIR"].iter_rows(min_row=2, values_only=True) if r[0] == "PPYRS")
    ws.append([de[0], "POA", de[2], None, "DE", "IMP_CLR", other[0], other[3]])
    newer = tmp_path / "m2c_newer.xlsx"
    wb.save(newer)
    res = loader.load(newer, db_url, actor="test")
    assert res["status"] == "LOADED" and res["batch_id"] != loaded["batch_id"]
    assert q(db_url, "SELECT batch_id FROM sirdb.import_batch WHERE is_current")[0][0] == res["batch_id"]
    assert q(db_url, "SELECT count(*) FROM sirdb.technical_sirg WHERE batch_id=%s", loaded["batch_id"])[0][0] == 1004
    with connect(db_url) as c, c.cursor() as cur:
        amb = resolve.resolve_sirg(cur, "DE", "POA")
        assert amb.status == "AMBIGUOUS" and amb.cmpt_nm is None
        old = resolve.resolve_sirg(cur, "DE", "POA", batch_id=loaded["batch_id"])
        assert old.status == "RESOLVED"
        v = resolve.validate_rule(cur, {"country": "DE", "sirg_cd": "POA", "effective_date": "01-01-2027"})
        assert not v["valid"] and any("different targets" in e["message"] for e in v["errors"])
    # Put the original back as current for the remaining tests.
    loader.load(workbook, db_url, actor="test")
    assert q(db_url, "SELECT batch_id FROM sirdb.import_batch WHERE is_current")[0][0] == loaded["batch_id"]


# --------------------------------------------------------------------------- #
# validation of proposed rules
# --------------------------------------------------------------------------- #

def test_valid_rule_has_only_warnings(db_url, loaded) -> None:
    with connect(db_url) as c, c.cursor() as cur:
        v = resolve.validate_rule(cur, {"country": "DE", "sirg_cd": "POA", "region": "EU", "shipment_type": "B2C",
                                        "effective_date": "15-11-2026", "currency": "EUR", "value_basis": "DV",
                                        "min_value": 150, "other_logic": 'incotermCodeList.contains("DDP")'})
    assert v["valid"] and v["resolution"]["cmpt_nm"] == "DEPOA"
    msgs = " ".join(w["message"] for w in v["warnings"])
    assert "not evaluated" in msgs and "rows [38]" in msgs


@pytest.mark.parametrize("rule,needle", [
    ({"country": "DE", "sirg_cd": "POA", "effective_date": "01-01-2027", "document_reference": "TEST|TEST|TEST"},
     "placeholder"),
    ({"country": "DE", "sirg_cd": "POA"}, "effective date is required"),
    ({"country": "DE", "sirg_cd": "POA", "effective_date": "01-01-2027", "expiry_date": "01-01-2026"},
     "before effective"),
    ({"country": "DE", "sirg_cd": "POA", "effective_date": "31-02-2027"}, "Effective date"),
    ({"country": "DE", "sirg_cd": "POA", "effective_date": "01-01-2027", "min_value": 500, "max_value": 100,
      "currency": "EUR", "value_basis": "DV"}, "above max_value"),
    ({"country": "DE", "sirg_cd": "POA", "effective_date": "01-01-2027", "min_weight": -1}, "0 or more"),
    ({"country": "DE", "sirg_cd": "POA", "effective_date": "01-01-2027", "min_value": 150}, "needs a currency"),
    ({"country": "DE", "sirg_cd": "XYZ", "effective_date": "01-01-2027"}, "maintenance"),
    ({"country": "Germany", "sirg_cd": "POA", "effective_date": "01-01-2027"}, "ISO two-letter"),
])
def test_blocking_errors(db_url, loaded, rule, needle) -> None:
    with connect(db_url) as c, c.cursor() as cur:
        v = resolve.validate_rule(cur, rule)
    assert not v["valid"]
    assert any(needle in e["message"] for e in v["errors"]), v["errors"]


# --------------------------------------------------------------------------- #
# final Import output
# --------------------------------------------------------------------------- #

def test_final_import_reproduces_the_workbook_exactly(db_url, loaded, workbook, tmp_path) -> None:
    out = tmp_path / "import.xlsx"
    assert resolve.export_import(out, db_url) == 56
    a = openpyxl.load_workbook(workbook, data_only=True)["Import"]
    b = openpyxl.load_workbook(out)["Import"]
    diffs = [(r, c) for r in range(1, a.max_row + 1) for c in range(1, 18)
             if a.cell(r, c).value != b.cell(r, c).value and not (a.cell(r, c).value in (None, "")
                                                                 and b.cell(r, c).value in (None, ""))]
    assert diffs == []


def test_approved_proposal_appears_in_the_final_import(db_url, loaded) -> None:
    rule = {"country": "BG", "sirg_cd": "POAB", "region": "EU", "shipment_type": "B2C", "effective_date": "01-12-2026",
            "network_product_code": "H,V,Y,M,E,P,8,J,3", "value_basis": "DV", "currency": "EUR",
            "min_value": 150.01, "other_logic": 'serviceCodeList.contains("DD")', "document_reference": "BG|Import|POAB"}
    bad = dict(rule, expiry_date="01-01-2026")
    with connect(db_url, actor="test") as c:
        saved_bad = resolve.save_proposal(c, bad, "tester")
        assert not saved_bad["valid"]
        with pytest.raises(resolve.ResolveError, match="blocking errors"):
            resolve.approve_proposal(c, saved_bad["proposal_id"], "approver")
        saved = resolve.save_proposal(c, rule, "tester")
        assert saved["valid"]
        resolve.approve_proposal(c, saved["proposal_id"], "approver")
        with c.cursor() as cur:
            rows = [r for r in resolve.final_import(cur) if r["origin"] == "PROPOSAL"]
    assert len(rows) == 1
    r = rows[0]
    assert r["SIRG Document Reference"] == "BG|Import|POAB"
    assert r["Effective Date (DD-MM-YYY)"] == "01-12-2026" and r["Min Value (>=)"] == "150.01"
    ops = q(db_url, "SELECT operation, new_row->>'status' FROM sirdb.change_log WHERE table_name='rule_proposal' "
                    "ORDER BY change_id")
    assert [o for o, _ in ops].count("INSERT") == 2 and ("UPDATE", "APPROVED") in ops


def test_profile_counts_match_the_known_workbook(workbook) -> None:
    rep = profile(read_workbook(workbook))
    inv = {s["sheet"]: s for s in rep["inventory"]}
    assert inv["Import"]["data_rows"] == 56 and inv["SIR Reference"]["data_rows"] == 85
    assert inv["Technical SIR"]["data_rows"] == 984 and inv["Technical SIRG"]["data_rows"] == 1004
    assert inv["Import"]["formula_cells"] == 33 and inv["SIR Reference"]["formula_cells"] == 34
    assert rep["summary"]["findings"].get("error", 0) == 0
