"""Ingest the M2C workbook.

Read once, to migrate off it. The workbook joins its tabs with pipe-delimited
strings produced by ``CONCATENATE``, and ``SIR Reference`` points back at
``Import`` by *row offset* -- so inserting a row silently re-points bindings.
Ingestion resolves those joins into real foreign keys so row order stops
mattering, and records the original key so the convention defect is reportable.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

import openpyxl
from sqlalchemy import delete
from sqlalchemy.orm import Session

from tools.core.models import (
    ExpressionTemplate, MirrorSir, MirrorSirg, PermittedValue, Rule, RuleBinding, RuleSet,
)

# e.g. "M2c_6_5_ver_1_W29_2.10.9-beta.2_20262608-131409_...xlsx"
_VERSION_RE = re.compile(r"_((?:\d+\.){2}\d+-[A-Za-z]+\.\d+)_")
_DATE_FORMATS = ("%d-%m-%Y", "%d-%m-%y", "%Y-%m-%d", "%d/%m/%Y")

SHEET_IMPORT = "Import"
SHEET_EXPORT = "Export"
SHEET_SIR_REF = "SIR Reference"
SHEET_TECH_SIR = "Technical SIR"
SHEET_TECH_SIRG = "Technical SIRG"
SHEET_OTHER_LOGIC = "Other Logic"
SHEET_VALIDATION = "Validation"

# Validation tab column header -> canonical field name
_PERMITTED_FIELDS = {
    "Region": "region",
    "B2B/B2C": "shipment_type",
    "Declared Value/ Intristic Value": "value_basis",
    "Currency List": "currency",
    "Waived if Whitelisted": "waiver",
    "Country": "country",
}


@dataclass
class WorkbookIngestResult:
    version: str
    source_name: str
    rule_set_id: str = ""
    rules: int = 0
    bindings: int = 0
    mirror_sir: int = 0
    mirror_sirg: int = 0
    permitted: int = 0
    templates: int = 0
    unbound_rules: list[str] = field(default_factory=list)
    orphan_bindings: list[str] = field(default_factory=list)


def detect_version(path: Path) -> str:
    match = _VERSION_RE.search(path.name)
    return match.group(1) if match else path.stem[:64]


def _s(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _num(value: Any) -> Optional[float]:
    text = _s(value).replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _date(value: Any) -> Optional[dt.date]:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _s(value)
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _sheet_dicts(workbook: Any, name: str) -> list[dict[str, Any]]:
    """Rows as dicts keyed by header. Blank headers become ``colN`` so the
    unnamed trailing column in ``SIR Reference`` is still reachable."""
    if name not in workbook.sheetnames:
        return []
    sheet = workbook[name]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []
    headers = [_s(h) or f"col{i}" for i, h in enumerate(header_row)]
    out: list[dict[str, Any]] = []
    for idx, row in enumerate(rows, start=2):
        if not any(cell is not None and _s(cell) for cell in row):
            continue
        record: dict[str, Any] = {"__row__": idx}
        for i, cell in enumerate(row):
            record[headers[i] if i < len(headers) else f"col{i}"] = cell
        out.append(record)
    return out


def _placeholder(value: str) -> bool:
    """Empty values and TEST|TEST|... placeholders never join anything. (A TBD hold code inside an
    otherwise real key, e.g. IL|Import|TBD|CIDM, still joins: both sides carry the same key.)"""
    return not value or any(p.strip().upper() == "TEST" for p in value.split("|"))


def _split_attribute_ids(raw: Any) -> list[str]:
    """``AttributeID`` holds one or more ids joined by ``&`` (sometimes ``&&``)."""
    text = _s(raw)
    if not text:
        return []
    return [part.strip() for part in text.replace("&&", "&").split("&") if part.strip()]


def ingest_workbook(
    session: Session,
    path: str | Path,
    version: Optional[str] = None,
    case_id: Optional[str] = None,
) -> WorkbookIngestResult:
    path = Path(path)
    version = version or detect_version(path)
    # data_only=True gives the cached CONCATENATE results, which are the join keys.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = WorkbookIngestResult(version=version, source_name=path.name)

    try:
        for model in (MirrorSir, MirrorSirg, PermittedValue, ExpressionTemplate):
            session.execute(delete(model).where(model.source_version == version))
        session.flush()

        # ---- mirror tabs (drift measurement only) ----
        for row in _sheet_dicts(workbook, SHEET_TECH_SIR):
            sir_id = _s(row.get("SIR Technical ID"))
            if not sir_id:
                continue
            session.add(MirrorSir(
                sir_id=sir_id,
                cmpt_nm=_s(row.get("CmptNm")) or None,
                nm=_s(row.get("Nm")) or None,
                dsc=_s(row.get("Dsc")) or None,
                source_version=version,
            ))
            result.mirror_sir += 1

        for row in _sheet_dicts(workbook, SHEET_TECH_SIRG):
            sirg_id = _s(row.get("SIRG Technical ID"))
            if not sirg_id:
                continue
            session.add(MirrorSirg(
                sirg_id=sirg_id,
                sirg_cd=_s(row.get("SIRG Cd")) or None,
                sirg_nm=_s(row.get("SIRG Nm")) or None,
                sirg_dsc=_s(row.get("SIRG Dsc")) or None,
                ctry_cd=_s(row.get("DHL Ctry Cd")) or None,
                proc_lvl_cd=_s(row.get("Processing Type")) or None,
                cmpt_nm=_s(row.get("CmptNm")) or None,
                sir_id=_s(row.get("SIR Technical ID")) or None,
                source_version=version,
            ))
            result.mirror_sirg += 1

        # ---- lookup lists and the expression palette ----
        for row in _sheet_dicts(workbook, SHEET_VALIDATION):
            for header, canonical in _PERMITTED_FIELDS.items():
                value = _s(row.get(header))
                if value:
                    session.add(PermittedValue(field=canonical, value=value, source_version=version))
                    result.permitted += 1

        for row in _sheet_dicts(workbook, SHEET_OTHER_LOGIC):
            expression = _s(row.get("Other Logic"))
            if expression:
                session.add(ExpressionTemplate(expression=expression, source_version=version))
                result.templates += 1

        # ---- rule rows ----
        rule_set = RuleSet(
            version="v0.1",
            label=f"Imported from {path.name}",
            case_id=case_id,
            created_by="ingest",
            source_version=version,
            sealed=False,
        )
        session.add(rule_set)
        session.flush()
        result.rule_set_id = rule_set.id

        # SIR Reference lines join an Import rule in one of two documented ways:
        #   legacy (hold-code rules): SIR Reference.Document Reference == Import.Document Reference
        #   SIRG rules: first three parts of SIR Reference.SIR Document Reference
        #               (Country|Import|SIRGCd|CmptNm) == Import.SIRG Document Reference
        # Placeholders (TEST...) never join.
        bindings_by_key: dict[str, list[dict[str, Any]]] = {}
        bindings_by_sirg: dict[str, list[dict[str, Any]]] = {}
        all_binding_rows: list[int] = []
        for row in _sheet_dicts(workbook, SHEET_SIR_REF):
            all_binding_rows.append(row["__row__"])
            key = _s(row.get("Document Reference"))
            if key and not _placeholder(key):
                bindings_by_key.setdefault(key, []).append(row)
            sdr = _s(row.get("SIR Document Reference")).split("|")
            if len(sdr) == 4 and not _placeholder("|".join(sdr)):
                bindings_by_sirg.setdefault("|".join(sdr[:3]), []).append(row)

        used_rows: set[int] = set()

        for sheet_name, direction in ((SHEET_IMPORT, "Import"), (SHEET_EXPORT, "Export")):
            for row in _sheet_dicts(workbook, sheet_name):
                ctry = _s(row.get("Country"))
                if not ctry:
                    continue
                doc_ref = _s(row.get("Document Reference"))
                sirg_ref = _s(row.get("SIRG Document Reference"))
                # "BG|Import|IOSS" style refs carry the SIRG code in slot 3.
                sirg_cd = ""
                parts = sirg_ref.split("|")
                if len(parts) == 3 and set(parts) != {"TEST"}:
                    sirg_cd = parts[2]

                rule = Rule(
                    rule_set_id=rule_set.id,
                    row_no=row["__row__"],
                    direction=direction,
                    region=_s(row.get("Region")) or None,
                    ctry_cd=ctry,
                    shipment_type=_s(row.get("B2B/B2C")) or None,
                    effective_date=_date(row.get("Effective Date (DD-MM-YYY)")),
                    expiry_date=_date(row.get("Date of Expiry (DD-MM-YYY)")),
                    network_product_codes=_s(row.get("Network Product Code")) or None,
                    value_basis=_s(row.get("Declared Value/ Intristic Value")) or None,
                    currency=_s(row.get("Currency")) or None,
                    min_value=_num(row.get("Min Value (>=)")),
                    max_value=_num(row.get("Max Value (<=)")),
                    min_weight=_num(row.get("Shipt Weight (kg) (>=)")),
                    max_weight=_num(row.get("Shipt Weight (kg) (<=)")),
                    other_logic=_s(row.get("Other Logic")) or None,
                    country_hold_cd=_s(row.get("Country Hold Code")) or None,
                    global_hold_cd=_s(row.get("Global Hold Code")) or None,
                    sirg_cd=sirg_cd or None,
                    doc_reference=doc_ref or None,
                )
                session.add(rule)
                session.flush()
                result.rules += 1

                matched, seen_rows = [], set()
                candidates = ([] if _placeholder(doc_ref) else bindings_by_key.get(doc_ref, [])) +                              ([] if _placeholder(sirg_ref) else bindings_by_sirg.get(sirg_ref, []))
                for b in candidates:
                    if b["__row__"] not in seen_rows:
                        seen_rows.add(b["__row__"])
                        matched.append(b)
                used_rows |= seen_rows
                if not matched:
                    result.unbound_rules.append(f"{ctry} row {row['__row__']} key={doc_ref or '(none)'}")
                for binding_row in matched:
                    session.add(RuleBinding(
                        rule_id=rule.id,
                        sirg_cd=sirg_cd or None,
                        sir_id=_s(binding_row.get("SIR Technical Id")) or None,
                        cmpt_nm=_s(binding_row.get("CmptNm")) or None,
                        waived_if_whitelisted=_s(binding_row.get("Waived if Whitelisted")).upper() == "Y",
                        waived_to_validate=_s(binding_row.get("Waived to Validate")).upper() == "Y",
                        check_service_code=_s(binding_row.get("Check if Service Code=")) or None,
                        attribute_ids=_split_attribute_ids(binding_row.get("AttributeID")),
                    ))
                    result.bindings += 1

        result.orphan_bindings = [f"SIR Reference row {r}" for r in all_binding_rows if r not in used_rows]
        session.flush()
    finally:
        workbook.close()

    return result
