"""Phase 1 — profile the M2C workbook: inventory, column profiles and reconciliation findings.

Read-only and deterministic: the same workbook always gives the same report. The
findings are reused by the loader, which stores them in ``sirdb.load_issue``.

Findings never change data. Each says what was seen, where (sheet, row, column),
how serious it is, and whether it is a verified fact or an inference.

Severity:
    error    the row cannot be used as it stands (e.g. a date that does not parse)
    warning  usable, but needs a business or data decision (placeholders, anomalies)
    info     a verified fact worth knowing (formula vs literal, unheaded column)

Run:  python -m tools.sirdb.profile <workbook.xlsx> [--out docs/sirdb/generated]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Optional

from .workbook import Sheet, Workbook, read_workbook

PLACEHOLDER = re.compile(r"\bTEST\b|\bTBD\b", re.I)
DOCREF_FORMULA = re.compile(r'^=CONCATENATE\(B(\d+),"\|Import\|",N(\d+),"\|",O(\d+)\)$')
SREF_FORMULA = re.compile(r'^=IF\(Import!P(\d+)<>0,Import!P(\d+),""\)$')
GUID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)
PROC_IMPORT = "IMP_CLR"     # every Technical SIRG row; "Import" in references is taken to mean this (see Q-03)

IMPORT_COLS = {
    "region": 1, "country": 2, "shipment_type": 3, "effective_date": 4, "expiry_date": 5,
    "network_product_code": 6, "value_basis": 7, "currency": 8, "min_value": 9, "max_value": 10,
    "min_weight": 11, "max_weight": 12, "other_logic": 13, "country_hold_code": 14,
    "global_hold_code": 15, "document_reference": 16, "sirg_document_reference": 17,
}
SREF_COLS = {"document_reference": 1, "sir_document_reference": 2, "waived_if_whitelisted": 3,
             "waived_to_validate": 4, "check_if_service_code": 5, "cmpt_nm": 6, "sir_technical_id": 7,
             "attribute_id": 8, "col_09_unheaded": 9}


@dataclass
class Finding:
    code: str
    severity: str           # error | warning | info
    sheet: str
    row: Optional[int]
    column: Optional[str]
    value: Any
    message: str
    basis: str = "verified"  # verified | inferred

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


# --------------------------------------------------------------------------- #
# value helpers (shared with the loader)
# --------------------------------------------------------------------------- #

def parse_dmy(value: Any) -> tuple[Optional[dt.date], Optional[str]]:
    """'01-01-2021' (text, DD-MM-YYYY) or a real date -> (date, None) or (None, reason)."""
    if value is None or value == "":
        return None, None
    if isinstance(value, (dt.date, dt.datetime)):
        return (value.date() if isinstance(value, dt.datetime) else value), None
    s = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(T.*)?", s):
        try:
            return dt.date.fromisoformat(s[:10]), None
        except ValueError:
            return None, "not a valid ISO date"
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", s)
    if not m:
        return None, "not DD-MM-YYYY"
    try:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))), None
    except ValueError:
        return None, "day or month out of range"


def parse_number(value: Any) -> tuple[Optional[float], Optional[str]]:
    if value is None or value == "":
        return None, None
    if isinstance(value, bool):
        return None, "boolean, not a number"
    if isinstance(value, (int, float)):
        return float(value), None
    try:
        return float(str(value).strip().replace(",", ".")), None
    except ValueError:
        return None, "not a number"


def split_ref(value: Any) -> list[str]:
    return str(value).split("|") if value not in (None, "") else []


def is_placeholder(value: Any) -> bool:
    return bool(value) and bool(PLACEHOLDER.search(str(value)))


# --------------------------------------------------------------------------- #
# profile
# --------------------------------------------------------------------------- #

def column_profiles(sheet: Sheet) -> list[dict[str, Any]]:
    out = []
    data_rows = sheet.data_rows()
    for col in range(1, sheet.max_col + 1):
        values, formulas, types = [], 0, Counter()
        for r in data_rows:
            for c in sheet.rows.get(r, []):
                if c.col == col:
                    values.append(c.value)
                    formulas += c.formula is not None
                    types[type(c.value).__name__] += 1
        distinct = Counter(str(v) for v in values)
        out.append({
            "column": col, "letter": _letter(col), "header": sheet.headers[col - 1],
            "non_empty": len(values), "empty": len(data_rows) - len(values), "formulas": formulas,
            "types": dict(types), "distinct": len(distinct),
            "top_values": distinct.most_common(5),
            "leading_or_trailing_space": sum(1 for v in values if isinstance(v, str) and v != v.strip()),
        })
    return out


def _letter(col: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(col)


def profile(wb: Workbook) -> dict[str, Any]:
    findings: list[Finding] = []
    add = lambda *a, **k: findings.append(Finding(*a, **k))  # noqa: E731

    imp, sref = wb["Import"], wb["SIR Reference"]
    tsir, tsirg = wb["Technical SIR"], wb["Technical SIRG"]
    val, other = wb["Validation"], wb["Other Logic"]

    # ---- reference sets -------------------------------------------------- #
    sir_by_id = {tsir.value(r, 4): (r, tsir.value(r, 1)) for r in tsir.data_rows()}
    sirg_key: dict[tuple, list[int]] = defaultdict(list)
    for r in tsirg.data_rows():
        sirg_key[(tsirg.value(r, 5), tsirg.value(r, 6), tsirg.value(r, 2))].append(r)
    lists = {val.headers[c - 1]: [val.value(r, c) for r in val.data_rows() if val.value(r, c) not in (None, "")]
             for c in range(1, val.max_col + 1)}
    vl = {k: {str(x).strip() for x in v} for k, v in lists.items()}
    import_p = {r: imp.value(r, 16) for r in imp.data_rows()}

    # ---- Import ------------------------------------------------------------ #
    last_filtered = int(re.findall(r"\d+", imp.autofilter or "A1:A0")[-1]) if imp.autofilter else None
    business_key = Counter()
    for r in imp.data_rows():
        g = lambda k: imp.value(r, IMPORT_COLS[k])  # noqa: E731
        country = g("country")
        eff, eff_err = parse_dmy(g("effective_date"))
        exp, exp_err = parse_dmy(g("expiry_date"))
        for k, d, err in (("effective_date", eff, eff_err), ("expiry_date", exp, exp_err)):
            raw = g(k)
            if err:
                add("IMP_DATE_INVALID", "error", "Import", r, k, raw, f"Date '{raw}' could not be read: {err}.")
            elif isinstance(raw, str):
                add("IMP_DATE_TEXT", "info", "Import", r, k, raw,
                    "Date stored as text (DD-MM-YYYY), not as an Excel date; normalised value kept separately.")
        if eff and exp and exp < eff:
            add("IMP_EXPIRY_BEFORE_EFFECTIVE", "warning", "Import", r, "expiry_date", g("expiry_date"),
                f"Expiry {exp:%d-%m-%Y} is before effective {eff:%d-%m-%Y}.")
        nums = {}
        for k in ("min_value", "max_value", "min_weight", "max_weight"):
            n, err = parse_number(g(k))
            nums[k] = n
            if err:
                add("IMP_NUMBER_INVALID", "error", "Import", r, k, g(k), f"'{g(k)}' is {err}.")
            elif n is not None and n < 0:
                add("IMP_NUMBER_NEGATIVE", "warning", "Import", r, k, g(k), "Negative threshold.")
        for lo, hi in (("min_value", "max_value"), ("min_weight", "max_weight")):
            if nums[lo] is not None and nums[hi] is not None and nums[lo] > nums[hi]:
                add("IMP_MIN_ABOVE_MAX", "warning", "Import", r, hi, g(hi), f"{lo} {nums[lo]} > {hi} {nums[hi]}.")
        for k, lst in (("region", "Region"), ("shipment_type", "B2B/B2C"),
                       ("value_basis", "Declared Value/ Intristic Value"), ("currency", "Currency List"),
                       ("country", "Country")):
            v = g(k)
            if v not in (None, "") and lst in vl and str(v).strip() not in vl[lst]:
                add("IMP_NOT_IN_VALIDATION_LIST", "warning", "Import", r, k, v,
                    f"'{v}' is not in the Validation tab list '{lst}' (the list is not enforced by Excel).")
        # Document Reference (P): formula or literal; convention Country|Import|CountryHold|GlobalHold
        p_val, p_formula = g("document_reference"), imp.formula(r, 16)
        expected = f"{country}|Import|{g('country_hold_code') or ''}|{g('global_hold_code') or ''}"
        if p_formula:
            m = DOCREF_FORMULA.match(p_formula)
            if not m or {int(x) for x in m.groups()} != {r}:
                add("IMP_DOCREF_FORMULA_UNEXPECTED", "warning", "Import", r, "document_reference", p_formula,
                    "Formula does not follow =CONCATENATE(B<n>,\"|Import|\",N<n>,\"|\",O<n>) for its own row.")
            else:
                add("IMP_DOCREF_FORMULA", "info", "Import", r, "document_reference", p_val,
                    "Document Reference is the CONCATENATE formula (Country|Import|CountryHold|GlobalHold).")
        else:
            add("IMP_DOCREF_LITERAL", "info", "Import", r, "document_reference", p_val,
                "Document Reference is a typed value, not the formula"
                + ("." if p_val != expected else " (it equals the formula result)."))
        if is_placeholder(p_val):
            add("IMP_DOCREF_PLACEHOLDER", "warning", "Import", r, "document_reference", p_val,
                "Document Reference is a placeholder (TEST).")
        if is_placeholder(g("country_hold_code")) or is_placeholder(g("global_hold_code")):
            add("IMP_HOLD_PLACEHOLDER", "warning", "Import", r, "country_hold_code", g("country_hold_code"),
                "Hold code is a placeholder (TBD).")
        if "/" in str(g("country_hold_code") or ""):
            add("IMP_HOLD_COMPOSITE", "warning", "Import", r, "country_hold_code", g("country_hold_code"),
                "Hold code holds two codes separated by '/'.")
        # SIRG Document Reference (Q): Country|Import|SIRGCd
        q = g("sirg_document_reference")
        parts = split_ref(q)
        if is_placeholder(q):
            add("IMP_SIRGREF_PLACEHOLDER", "warning", "Import", r, "sirg_document_reference", q,
                "SIRG Document Reference is a placeholder (TEST).")
        elif len(parts) != 3 or parts[1] != "Import":
            add("IMP_SIRGREF_FORMAT", "warning", "Import", r, "sirg_document_reference", q,
                "Expected Country|Import|SIRGCd.")
        else:
            if parts[0] != country:
                add("IMP_SIRGREF_COUNTRY", "warning", "Import", r, "sirg_document_reference", q,
                    f"Country in the reference ({parts[0]}) differs from the row's country ({country}).")
            hits = sirg_key.get((parts[0], PROC_IMPORT, parts[2]), [])
            if not hits:
                add("IMP_SIRGREF_UNRESOLVED", "warning", "Import", r, "sirg_document_reference", q,
                    f"No Technical SIRG row for ({parts[0]}, {PROC_IMPORT}, {parts[2]}).")
            elif len(hits) > 1:
                add("IMP_SIRGREF_AMBIGUOUS", "warning", "Import", r, "sirg_document_reference", q,
                    f"{len(hits)} Technical SIRG rows match (rows {hits}).")
        has_business = any(g(k) not in (None, "") for k in ("other_logic", "min_value", "max_value",
                                                             "min_weight", "max_weight"))
        if has_business and is_placeholder(p_val) and is_placeholder(q):
            add("IMP_BUSINESS_WITHOUT_REFERENCE", "warning", "Import", r, None, None,
                "Business conditions are filled in but both references are placeholders.")
        if last_filtered and r > last_filtered:
            add("IMP_OUTSIDE_AUTOFILTER", "info", "Import", r, None, imp.autofilter,
                f"Row is outside the sheet's filter range {imp.autofilter} (likely appended later).", basis="inferred")
        business_key[tuple(str(g(k)) for k in IMPORT_COLS if k != "document_reference")] += 1

    dupes = [k for k, n in business_key.items() if n > 1]
    if dupes:
        add("IMP_DUPLICATE_CANDIDATE", "warning", "Import", None, None, len(dupes),
            f"{len(dupes)} group(s) of rows identical in every column except Document Reference.")

    # ---- SIR Reference ------------------------------------------------------- #
    import_p_values = set(v for v in import_p.values() if v)
    import_q_values = {imp.value(r, 17) for r in imp.data_rows()}
    for r in sref.data_rows():
        g = lambda k: sref.value(r, SREF_COLS[k])  # noqa: E731
        a, a_formula, b = g("document_reference"), sref.formula(r, 1), g("sir_document_reference")
        if a_formula:
            m = SREF_FORMULA.match(a_formula)
            if not m or m.group(1) != m.group(2):
                add("SREF_FORMULA_UNEXPECTED", "warning", "SIR Reference", r, "document_reference", a_formula,
                    "Formula does not follow =IF(Import!P<n><>0,Import!P<n>,\"\").")
            else:
                tgt = int(m.group(1))
                if import_p.get(tgt) != a:
                    add("SREF_FORMULA_STALE", "warning", "SIR Reference", r, "document_reference", a,
                        f"Cached value differs from Import!P{tgt} ({import_p.get(tgt)}).")
                add("SREF_FORMULA_POSITIONAL", "info", "SIR Reference", r, "document_reference", a_formula,
                    f"Links to Import row {tgt} by position: inserting or deleting Import rows re-points it.")
        if a not in (None, "") and not is_placeholder(a) and a not in import_p_values:
            add("SREF_DOCREF_ORPHAN", "warning", "SIR Reference", r, "document_reference", a,
                "No Import row has this Document Reference.")
        bparts = split_ref(b)
        if is_placeholder(b):
            pass
        elif len(bparts) != 4 or bparts[1] != "Import":
            add("SREF_SIRDOCREF_FORMAT", "warning", "SIR Reference", r, "sir_document_reference", b,
                "Expected Country|Import|SIRGCd|CmptNm.")
        else:
            hits = sirg_key.get((bparts[0], PROC_IMPORT, bparts[2]), [])
            if not hits:
                add("SREF_SIRG_UNRESOLVED", "warning", "SIR Reference", r, "sir_document_reference", b,
                    f"No Technical SIRG row for ({bparts[0]}, {PROC_IMPORT}, {bparts[2]}).")
            elif all(tsirg.value(h, 7) != bparts[3] for h in hits):
                add("SREF_SIRG_COMPONENT_MISMATCH", "warning", "SIR Reference", r, "sir_document_reference", b,
                    f"Technical SIRG gives component {[tsirg.value(h, 7) for h in hits]}, the reference says {bparts[3]}.")
            if "|".join(bparts[:3]) not in import_q_values:
                add("SREF_NO_IMPORT_RULE", "info", "SIR Reference", r, "sir_document_reference", b,
                    "No Import row has the SIRG Document Reference " + "|".join(bparts[:3]) + ".")
        sid, cm = g("sir_technical_id"), g("cmpt_nm")
        if sid:
            if sid not in sir_by_id:
                add("SREF_SIR_ID_UNKNOWN", "warning", "SIR Reference", r, "sir_technical_id", sid,
                    "SIR Technical Id is not in Technical SIR.")
            elif sir_by_id[sid][1] != cm:
                add("SREF_SIR_COMPONENT_MISMATCH", "warning", "SIR Reference", r, "cmpt_nm", cm,
                    f"Technical SIR names this id '{sir_by_id[sid][1]}'.")
        for k in ("waived_if_whitelisted", "waived_to_validate"):
            v = g(k)
            if v not in (None, "") and str(v).strip() not in ("Y", "N"):
                add("SREF_FLAG_VALUE", "warning", "SIR Reference", r, k, v, f"'{v}' is not Y or N.")
        ids = [x for x in re.split(r"&+", str(g("attribute_id") or "")) if x]
        labels = [x for x in re.split(r"&+", str(g("col_09_unheaded") or "")) if x.strip()]
        if ids and labels and len(ids) != len(labels):
            add("SREF_ATTR_LABEL_COUNT", "warning", "SIR Reference", r, "col_09_unheaded", g("col_09_unheaded"),
                f"{len(ids)} attribute id(s) but {len(labels)} label(s) in the unheaded column.", basis="inferred")
        if "&&" in str(g("attribute_id") or "") or ("&" in str(g("attribute_id") or "")):
            seps = set(re.findall(r"&+", str(g("attribute_id"))))
            if seps - {"&&"}:
                add("SREF_ATTR_SEPARATOR", "info", "SIR Reference", r, "attribute_id", g("attribute_id"),
                    f"Attribute ids separated by {sorted(seps)} (most rows use '&&').")
    if sref.headers[8] is None:
        add("SREF_UNHEADED_COLUMN", "info", "SIR Reference", 1, "I", None,
            "Column I has no header; its values line up one-to-one with the AttributeID entries and read like "
            "attribute names. Stored as col_09_unheaded until confirmed (Q-01).", basis="inferred")

    # ---- Technical SIR / SIRG ---------------------------------------------- #
    for cm, n in Counter(tsir.value(r, 1) for r in tsir.data_rows()).items():
        if n > 1:
            add("TSIR_CMPT_REPEATED", "info", "Technical SIR", None, "CmptNm", cm,
                f"CmptNm '{cm}' appears on {n} rows with different SIR Technical IDs — CmptNm is not a key.")
    seen: dict[tuple, int] = {}
    for r in tsirg.data_rows():
        rec = tuple(tsirg.record(r))
        if rec in seen:
            add("TSIRG_EXACT_DUPLICATE", "warning", "Technical SIRG", r, None, rec[1],
                f"Identical to row {seen[rec]}.")
        else:
            seen[rec] = r
        sid = tsirg.value(r, 8)
        if sid not in sir_by_id:
            add("TSIRG_SIR_UNKNOWN", "error", "Technical SIRG", r, "SIR Technical ID", sid,
                "SIR Technical ID not in Technical SIR.")
        elif sir_by_id[sid][1] != tsirg.value(r, 7):
            add("TSIRG_COMPONENT_MISMATCH", "warning", "Technical SIRG", r, "CmptNm", tsirg.value(r, 7),
                f"Technical SIR names this id '{sir_by_id[sid][1]}'.")
    for key, rows_ in sirg_key.items():
        recs = {tuple(tsirg.record(x)) for x in rows_}
        if len(rows_) > 1 and len(recs) > 1:
            add("TSIRG_KEY_AMBIGUOUS", "warning", "Technical SIRG", rows_[0], None, "|".join(map(str, key)),
                f"(country, processing type, SIRG code) maps to {len(recs)} different rows: {rows_}.")

    # ---- Sheet1 / Sheet2 / Validation / Other Logic / Export ----------------- #
    s1 = wb["Sheet1"]
    for r in s1.data_rows():
        rec = s1.record(r)
        hits = sirg_key.get((rec[4], rec[5], rec[1]), []) if len(rec) >= 8 else []
        same = [h for h in hits if tsirg.record(h) == rec]
        if same:
            add("SHEET1_MATCHES_SIRG", "info", "Sheet1", r, None, rec[1], f"Identical to Technical SIRG row {same}.",
                basis="inferred")
        elif hits:
            diff = [tsirg.headers[i] for i in range(8) if tsirg.record(hits[0])[i] != rec[i]]
            add("SHEET1_DIFFERS_FROM_SIRG", "warning", "Sheet1", r, None, rec[1],
                f"Shaped like a Technical SIRG row; same key as row {hits[0]} but differs in {diff}.", basis="inferred")
    s2 = [wb["Sheet2"].value(r, 1) for r in wb["Sheet2"].data_rows()]
    legacy = {imp.value(r, 2) for r in imp.data_rows() if imp.formula(r, 16)}
    add("SHEET2_COUNTRY_LIST", "info", "Sheet2", None, "A", ",".join(map(str, s2)),
        f"{len(s2)} country codes, no header. Not in Validation.Country: {sorted(set(s2) - vl.get('Country', set()))}. "
        f"Countries of formula-based Import rows not in Sheet2: {sorted(legacy - set(s2))}.", basis="inferred")
    for c in range(1, val.max_col + 1):
        for r in val.data_rows():
            v = val.value(r, c)
            if isinstance(v, str) and v != v.strip():
                add("VAL_WHITESPACE", "warning", "Validation", r, val.headers[c - 1], v,
                    f"'{v}' has leading/trailing spaces.")
    if not any(s.data_validations for s in wb.sheets.values()):
        add("VAL_NOT_ENFORCED", "info", "Validation", None, None, None,
            "The workbook has no Excel data-validation rules: the Validation lists are reference only.")
    used = {imp.value(r, 13) for r in imp.data_rows() if imp.value(r, 13)}
    for r in other.data_rows():
        v = other.value(r, 1)
        exact = v in used
        add("OTHERLOGIC_ENTRY", "info", "Other Logic", r, "A", v,
            ("Used verbatim in Import." if exact else "Not used verbatim in Import.")
            + (" Has leading/trailing spaces." if isinstance(v, str) and v != v.strip() else "")
            + (" Uses ',' inside matchCiMask where Import uses '|'." if "matchCiMask(\"shortCode," in str(v) else ""))
    exp = wb["Export"]
    extra = [h for h in exp.headers if h not in imp.headers]
    add("EXPORT_HEADER_ONLY", "info", "Export", 1, None, exp.dimensions,
        f"Header row only ({len(exp.headers)} columns); compared with Import it adds {extra}. "
        f"Filter range {exp.autofilter} covers empty rows.", basis="inferred")

    inventory = [{
        "sheet": s.name, "index": s.index, "dimensions": s.dimensions, "max_row": s.max_row, "max_col": s.max_col,
        "headerless": s.headerless, "data_rows": len(s.data_rows()), "headers": s.headers, "autofilter": s.autofilter,
        "formula_cells": sum(1 for cells in s.rows.values() for c in cells if c.formula),
        "data_validations": s.data_validations, "columns": column_profiles(s),
    } for s in wb.sheets.values()]
    counts = Counter(f.severity for f in findings)
    return {"workbook": {"path": wb.path.name, "sha256": wb.sha256, "size": wb.size},
            "inventory": inventory, "findings": [f.as_dict() for f in findings],
            "summary": {"findings": dict(counts), "by_code": dict(Counter(f.code for f in findings))}}


# --------------------------------------------------------------------------- #
# report
# --------------------------------------------------------------------------- #

def to_markdown(rep: dict[str, Any]) -> str:
    w = rep["workbook"]
    out = [f"# Workbook profile and reconciliation (generated)\n",
           f"- File: `{Path(w['path']).name}`", f"- SHA-256: `{w['sha256']}`", f"- Size: {w['size']} bytes\n",
           "## Inventory\n", "| Sheet | Used range | Data rows | Columns | Formula cells | Header row | Filter |",
           "|---|---|---|---|---|---|---|"]
    for s in rep["inventory"]:
        out.append(f"| {s['sheet']} | {s['dimensions']} | {s['data_rows']} | {s['max_col']} | {s['formula_cells']} | "
                   f"{'none (data starts row 1)' if s['headerless'] else 'row 1'} | {s['autofilter'] or '—'} |")
    out.append("\n## Columns\n")
    for s in rep["inventory"]:
        out += [f"### {s['sheet']}\n", "| Col | Header | Filled | Empty | Formulas | Types | Distinct | Top values |",
                "|---|---|---|---|---|---|---|---|"]
        for c in s["columns"]:
            top = "; ".join(f"`{str(v)[:40]}` ×{n}" for v, n in c["top_values"][:3])
            out.append(f"| {c['letter']} | {c['header'] if c['header'] is not None else '*(no header)*'} | "
                       f"{c['non_empty']} | {c['empty']} | {c['formulas']} | {c['types']} | {c['distinct']} | {top} |")
        out.append("")
    out += ["## Findings by code\n", "| Code | Count |", "|---|---|"]
    out += [f"| {k} | {v} |" for k, v in sorted(rep["summary"]["by_code"].items())]
    out += ["\n## Every finding (row-level reconciliation)\n",
            "| Severity | Code | Sheet | Row | Column | Value | Finding | Basis |", "|---|---|---|---|---|---|---|---|"]
    order = {"error": 0, "warning": 1, "info": 2}
    for f in sorted(rep["findings"], key=lambda f: (order[f["severity"]], f["sheet"], f["row"] or 0, f["code"])):
        v = str(f["value"])[:60].replace("|", "\\|") if f["value"] is not None else ""
        out.append(f"| {f['severity']} | {f['code']} | {f['sheet']} | {f['row'] or ''} | {f['column'] or ''} | "
                   f"{v} | {f['message'].replace('|', '\\|')} | {f['basis']} |")
    return "\n".join(out) + "\n"


def main(argv: Optional[list[str]] = None) -> None:
    ap = argparse.ArgumentParser(description="Profile the M2C workbook (read-only).")
    ap.add_argument("workbook")
    ap.add_argument("--out", default="docs/sirdb/generated")
    args = ap.parse_args(argv)
    rep = profile(read_workbook(args.workbook))
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "profile.json").write_text(json.dumps(rep, indent=2, default=str), encoding="utf-8")
    (out / "reconciliation.md").write_text(to_markdown(rep), encoding="utf-8")
    print(json.dumps({"sha256": rep["workbook"]["sha256"], **rep["summary"]}, indent=2))


if __name__ == "__main__":
    main()
